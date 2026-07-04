# DOCX 解析与入库。

from __future__ import annotations

import io


def extract_text_from_docx(data: bytes) -> str:
    """从 DOCX 字节提取纯文本。"""
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "python-docx 未安装。请执行: pip install python-docx 或 pip install -e '.[docx]'"
        ) from exc

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)
