# =============================================================================
# 实验数据多格式解析：JSON / CSV / TSV / Excel → summary + metrics。
# =============================================================================

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from typing import Any


@dataclass
class ParsedExperimentData:
    """统一解析结果，供入库与 DataPacket 使用。"""

    summary: dict[str, Any] = field(default_factory=dict)
    metrics: dict[str, Any] = field(default_factory=dict)
    format: str = "unknown"
    preview_rows: list[dict[str, Any]] = field(default_factory=list)


def _is_number(value: Any) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    if isinstance(value, str):
        try:
            float(value.strip())
            return True
        except ValueError:
            return False
    return False


def _to_number(value: Any) -> int | float | Any:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        s = value.strip()
        try:
            if re.fullmatch(r"-?\d+", s):
                return int(s)
            return float(s)
        except ValueError:
            return value
    return value


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _collect_notable_cells(rows: list[list[Any]], *, max_cells: int = 40) -> list[dict[str, Any]]:
    """宽表/脏表头时，抽出前若干非空单元格，便于 Agent 读数。"""
    notable: list[dict[str, Any]] = []
    for ri, row in enumerate(rows[:40]):
        for ci, raw in enumerate(row):
            text = _cell_str(raw)
            if not text:
                continue
            entry: dict[str, Any] = {"row": ri, "col": ci, "value": _to_number(raw) if _is_number(raw) else text}
            notable.append(entry)
            if len(notable) >= max_cells:
                return notable
    return notable


def _looks_like_placeholder_header(header: list[str]) -> bool:
    """首行多为空/占位 col_*，或重复空名。"""
    if not header:
        return True
    placeholders = sum(1 for h in header if not h or h.startswith("col_"))
    return placeholders >= max(1, len(header) // 2)


def _promote_header_from_label_column(cleaned: list[list[Any]]) -> tuple[list[str], list[list[Any]]] | None:
    """
    转置式宽表：首列是指标名，其余列为条件/模型取值。
    例：
      , Adam, SGD
      loss, 0.1, 0.2
      acc, 0.9, 0.8
    → metrics: Adam__loss=0.1, SGD__loss=0.2, ...
    """
    if len(cleaned) < 2:
        return None
    # 找「首列有较多非空文本标签、其它列多为数值」的数据区
    # 允许首行是列标题（首格可空）
    header_row = cleaned[0]
    col_labels = [_cell_str(c) or f"col_{i}" for i, c in enumerate(header_row)]
    # 若首列标题空，用 "metric"
    if col_labels[0].startswith("col_"):
        col_labels[0] = "metric"

    body = cleaned[1:]
    label_count = 0
    numeric_hits = 0
    numeric_total = 0
    for r in body:
        if not r:
            continue
        left = _cell_str(r[0])
        if left and not _is_number(left):
            label_count += 1
        for c in r[1:]:
            if _cell_str(c) == "":
                continue
            numeric_total += 1
            if _is_number(c):
                numeric_hits += 1

    if label_count < 2:
        return None
    if numeric_total == 0 or numeric_hits / max(1, numeric_total) < 0.5:
        return None

    # 合成「伪表头」：metric + 各列名；数据行保持
    return col_labels, body


def _metrics_from_label_wide_table(
    col_labels: list[str],
    body: list[list[Any]],
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    """首列标签 × 列标题 → 可读 metrics 键。"""
    metrics: dict[str, Any] = {}
    preview: list[dict[str, Any]] = []
    multi_col = len(col_labels) > 2
    for r in body:
        row_label = _cell_str(r[0]) if r else ""
        if not row_label:
            continue
        row_obj: dict[str, Any] = {"metric": row_label}
        for i, col_name in enumerate(col_labels[1:], start=1):
            raw = r[i] if i < len(r) else ""
            if _cell_str(raw) == "":
                continue
            val = _to_number(raw) if _is_number(raw) else _cell_str(raw)
            if col_name and not col_name.startswith("col_"):
                key = f"{col_name}__{row_label}" if multi_col else f"{col_name}__{row_label}"
            else:
                key = row_label if not multi_col else f"c{i}__{row_label}"
            if key in metrics:
                key = f"{key}__c{i}"
            metrics[key] = val
            row_obj[col_name] = val
        preview.append(row_obj)
    return metrics, preview[:20]


def _readable_summary_text(summary: dict[str, Any], metrics: dict[str, Any]) -> str:
    parts: list[str] = []
    shape = summary.get("shape")
    if shape:
        parts.append(f"shape={shape}")
    if summary.get("row_count") is not None:
        parts.append(f"rows={summary['row_count']}")
    cols = summary.get("columns")
    if isinstance(cols, list) and cols:
        shown = [str(c) for c in cols[:12]]
        parts.append("columns=" + ",".join(shown))
    if metrics:
        kv = [f"{k}={v}" for k, v in list(metrics.items())[:12]]
        parts.append("metrics: " + "; ".join(kv))
    notable = summary.get("notable_cells")
    if isinstance(notable, list) and notable and not metrics:
        cells = []
        for n in notable[:15]:
            if not isinstance(n, dict):
                continue
            cells.append(f"[{n.get('row')},{n.get('col')}]={n.get('value')}")
        if cells:
            parts.append("notable: " + "; ".join(cells))
    return " | ".join(parts) if parts else "no readable metrics"


def _rows_to_parsed(rows: list[list[Any]], *, fmt: str) -> ParsedExperimentData:
    """
    表格约定：
    1) 两列 key-value → 全部进 metrics
    2) 首行表头 + 多行数据 → 末行数值列进 metrics；summary 含列名/行数；preview 最多 20 行
    3) 宽表/多级脏表头 → notable_cells + 尽量抽出「列标题×行标签」可读 metrics
    """
    cleaned: list[list[Any]] = []
    for row in rows:
        if row is None:
            continue
        cells = [("" if c is None else c) for c in row]
        if all(str(c).strip() == "" for c in cells):
            continue
        cleaned.append(cells)

    if not cleaned:
        return ParsedExperimentData(
            summary={"note": "empty_table", "readable": "empty table"},
            metrics={},
            format=fmt,
        )

    notable = _collect_notable_cells(cleaned)

    # 两列：优先识别「表头+数据」；否则视为键值对（可带 metric/value 表头）
    if all(len(r) >= 2 for r in cleaned) and max(len(r) for r in cleaned) <= 2:
        first0, first1 = cleaned[0][0], cleaned[0][1]
        headerish = len(cleaned) >= 2 and not _is_number(first0) and not _is_number(first1)
        # 首列后续为数值 → 普通两列表格（如 lr,loss）
        if headerish and any(_is_number(r[0]) for r in cleaned[1:]):
            pass  # fall through to table path
        else:
            kv_rows = cleaned[1:] if headerish else cleaned
            metrics: dict[str, Any] = {}
            for r in kv_rows:
                key = str(r[0]).strip()
                if not key:
                    continue
                metrics[key] = _to_number(r[1])
            summary = {
                "shape": "key_value",
                "row_count": len(metrics),
                "notable_cells": notable[:20],
            }
            summary["readable"] = _readable_summary_text(summary, metrics)
            return ParsedExperimentData(
                summary=summary,
                metrics=metrics,
                format=fmt,
                preview_rows=[{"key": k, "value": v} for k, v in list(metrics.items())[:20]],
            )

    # 宽表：首列标签 + 多列数值
    wide = _promote_header_from_label_column(cleaned)
    if wide is not None:
        col_labels, body = wide
        metrics, preview = _metrics_from_label_wide_table(col_labels, body)
        if metrics:
            summary = {
                "shape": "wide_labeled",
                "columns": col_labels,
                "row_count": len(body),
                "notable_cells": notable[:20],
            }
            summary["readable"] = _readable_summary_text(summary, metrics)
            return ParsedExperimentData(
                summary=summary,
                metrics=metrics,
                format=fmt,
                preview_rows=preview,
            )

    header = [str(c).strip() or f"col_{i}" for i, c in enumerate(cleaned[0])]
    data_rows = cleaned[1:] if len(cleaned) > 1 else []

    # 脏表头：尝试用下一行非空更多的行当表头
    if _looks_like_placeholder_header(header) and data_rows:
        best_i = 0
        best_score = sum(1 for c in cleaned[0] if _cell_str(c))
        for i, r in enumerate(cleaned[1:6], start=1):
            score = sum(1 for c in r if _cell_str(c) and not _is_number(c))
            if score > best_score:
                best_score = score
                best_i = i
        if best_i > 0:
            header = [str(c).strip() or f"col_{i}" for i, c in enumerate(cleaned[best_i])]
            data_rows = cleaned[best_i + 1 :]

    dict_rows: list[dict[str, Any]] = []
    for r in data_rows:
        item: dict[str, Any] = {}
        for i, name in enumerate(header):
            raw = r[i] if i < len(r) else ""
            item[name] = _to_number(raw) if _is_number(raw) else (str(raw).strip() if raw is not None else "")
        dict_rows.append(item)

    metrics = {}
    if dict_rows:
        last = dict_rows[-1]
        for k, v in last.items():
            if _is_number(v) or isinstance(v, (int, float)):
                metrics[k] = _to_number(v)

    # 若 metrics 几乎全是 col_*，把 notable 提升为可读线索，并从预览抽命名字段
    if metrics and _looks_like_placeholder_header(list(metrics.keys())):
        named: dict[str, Any] = {}
        for row in dict_rows[:30]:
            # 找一行里的文本标签 + 邻近数值
            labels = [str(v) for v in row.values() if isinstance(v, str) and v and not _is_number(v)]
            nums = [(k, v) for k, v in row.items() if _is_number(v) or isinstance(v, (int, float))]
            if labels and nums:
                label = labels[0][:64]
                for k, v in nums[:4]:
                    key = f"{label}__{k}" if not str(k).startswith("col_") else label
                    if key in named:
                        key = f"{key}__{k}"
                    named[key] = _to_number(v)
        if named:
            metrics = {**named, **{k: v for k, v in metrics.items() if not str(k).startswith("col_")}}

    summary = {
        "shape": "table",
        "columns": header,
        "row_count": len(dict_rows),
        "notable_cells": notable[:20],
    }
    summary["readable"] = _readable_summary_text(summary, metrics)
    return ParsedExperimentData(
        summary=summary,
        metrics=metrics,
        format=fmt,
        preview_rows=dict_rows[:20],
    )


def parse_json_bytes(data: bytes) -> ParsedExperimentData:
    obj = json.loads(data.decode("utf-8-sig"))
    if isinstance(obj, list):
        # JSON 数组：当作记录表
        if obj and isinstance(obj[0], dict):
            header = list(obj[0].keys())
            rows = [header] + [[rec.get(h, "") for h in header] for rec in obj if isinstance(rec, dict)]
            parsed = _rows_to_parsed(rows, fmt="json")
            return parsed
        raise ValueError("JSON 数组须为对象列表")
    if not isinstance(obj, dict):
        raise ValueError("JSON 须为对象或对象数组")
    summary = obj.get("summary") if isinstance(obj.get("summary"), dict) else {}
    metrics = obj.get("metrics") if isinstance(obj.get("metrics"), dict) else {}
    if not summary and not metrics:
        # 整个对象当 metrics（数值字段）+ 其余进 summary
        metrics = {k: v for k, v in obj.items() if _is_number(v) or isinstance(v, (int, float))}
        summary = {k: v for k, v in obj.items() if k not in metrics}
        if not metrics and obj:
            metrics = obj
            summary = {"shape": "flat_object"}
    summary = dict(summary)
    metrics = dict(metrics)
    if "readable" not in summary:
        summary["readable"] = _readable_summary_text(summary, metrics)
    return ParsedExperimentData(summary=summary, metrics=metrics, format="json")


def parse_csv_bytes(data: bytes, *, delimiter: str = ",") -> ParsedExperimentData:
    text = data.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [list(r) for r in reader]
    fmt = "tsv" if delimiter == "\t" else "csv"
    parsed = _rows_to_parsed(rows, fmt=fmt)
    if parsed.preview_rows:
        parsed.summary = {**parsed.summary, "preview": parsed.preview_rows}
    return parsed


def parse_excel_bytes(data: bytes) -> ParsedExperimentData:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("服务端未安装 openpyxl，无法解析 Excel。请 pip install openpyxl") from exc

    wb = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows: list[list[Any]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
    finally:
        wb.close()
    parsed = _rows_to_parsed(rows, fmt="xlsx")
    if parsed.preview_rows:
        parsed.summary = {**parsed.summary, "preview": parsed.preview_rows}
    return parsed


def parse_experiment_file(filename: str, data: bytes) -> ParsedExperimentData:
    """按扩展名分发解析。"""
    name = (filename or "").lower().strip()
    if name.endswith(".json"):
        return parse_json_bytes(data)
    if name.endswith(".csv"):
        return parse_csv_bytes(data, delimiter=",")
    if name.endswith(".tsv"):
        return parse_csv_bytes(data, delimiter="\t")
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return parse_excel_bytes(data)
    if name.endswith(".xls"):
        raise ValueError("暂不支持旧版 .xls，请另存为 .xlsx 或导出 CSV")
    # 尝试按内容嗅探
    head = data[:64].lstrip()
    if head.startswith(b"{") or head.startswith(b"["):
        return parse_json_bytes(data)
    # 默认当 CSV
    return parse_csv_bytes(data, delimiter=",")
