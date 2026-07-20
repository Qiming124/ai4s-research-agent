# =============================================================================
# DOCX 导出：Markdown → pandoc → DOCX，回退 python-docx。
#
# 职责：
#     1. build_docx_bytes() 主入口，先 math_docx_prep 预处理公式
#     2. pandoc 可用时优先 convert_markdown_to_docx
#     3. 失败时回退纯 python-docx 段落渲染
#
# 架构位置：
#     - 被调用：server/api/export.py
#     - 调用：export/math_docx_prep.py、pandoc_converter.py
#
# 阅读提示：
#     - 新人先看 build_docx_bytes
#
# Debug：
#     - Word 公式空白 → prepare_markdown_for_docx 清洗 aligned / \\square
#     - pandoc 报错 → apt install pandoc 或检查临时目录权限
# =============================================================================

from __future__ import annotations

import io
import logging
import re
from typing import Any

logger = logging.getLogger(__name__)


def build_docx_bytes(*, markdown: str) -> bytes:
    """从 Markdown 生成 DOCX。"""
    from server.export.math_docx_prep import prepare_markdown_for_docx
    from server.export.pandoc_converter import convert_markdown_to_docx, pandoc_available

    # Word/WPS 对 OMML 中的中文、\\square、aligned 较敏感，先清洗再转
    markdown = prepare_markdown_for_docx(markdown)

    if pandoc_available():
        try:
            return convert_markdown_to_docx(markdown)
        except RuntimeError as exc:
            logger.warning("pandoc Markdown→DOCX 失败，回退 python-docx: %s", exc)

    return _docx_from_markdown_fallback(markdown)


def build_docx_bytes_from_entries(*, title: str, entries: list[dict[str, Any]]) -> bytes:
    """兼容旧调用：先转 Markdown 再导出 DOCX。"""
    from server.export.builder import build_markdown_document

    return build_docx_bytes(markdown=build_markdown_document(title=title, entries=entries))


def _docx_from_markdown_fallback(markdown: str) -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
        from docx.oxml.ns import qn
        from docx.shared import Pt, Cm
    except ImportError as exc:
        raise RuntimeError(
            "python-docx 未安装。请执行: pip install python-docx 或 pip install -e '.[docx]'"
        ) from exc

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    styles = doc.styles
    if "Normal" in styles:
        styles["Normal"].font.size = Pt(10.5)
        styles["Normal"].font.name = "Times New Roman"
        styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for style_name, size in (("Title", 16), ("Heading 1", 12), ("Heading 2", 11)):
        if style_name in styles:
            styles[style_name].font.size = Pt(size)
            styles[style_name].font.bold = True

    lines = markdown.splitlines()
    idx = 0
    in_abstract = False
    while idx < len(lines):
        line = lines[idx]
        if line.startswith("# "):
            heading = doc.add_heading(line[2:].strip(), level=0)
            heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            for run in heading.runs:
                run.font.size = Pt(16)
            in_abstract = False
            idx += 1
            continue
        if line.startswith("## "):
            title = line[3:].strip()
            h = doc.add_heading(title, level=1)
            for run in h.runs:
                run.font.size = Pt(12)
            low = title.lower()
            in_abstract = low.startswith("摘要") or low.startswith("abstract")
            idx += 1
            continue
        if line.startswith("### "):
            h = doc.add_heading(line[4:].strip(), level=2)
            for run in h.runs:
                run.font.size = Pt(11)
            in_abstract = False
            idx += 1
            continue
        if not line.strip():
            idx += 1
            continue
        para_lines = [line]
        idx += 1
        while idx < len(lines) and lines[idx].strip() and not lines[idx].startswith("#"):
            para_lines.append(lines[idx])
            idx += 1
        from server.export.html_pdf_renderer import _normalize_math_markup

        p = doc.add_paragraph(_normalize_math_markup("\n".join(para_lines)))
        for run in p.runs:
            run.font.size = Pt(9.5 if in_abstract else 10.5)

    if not doc.paragraphs:
        doc.add_paragraph(markdown or "（空文档）")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _kind_label(kind: str) -> str:
    labels = {
        "theorem": "定理",
        "lemma": "引理",
        "hypothesis": "假设",
        "note": "笔记",
        "definition": "定义",
    }
    return labels.get(kind, kind)
