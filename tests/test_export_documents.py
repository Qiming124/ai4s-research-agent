# 导出与文档解析测试。

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test-key-for-ci-only")
    from server.config import get_settings

    get_settings.cache_clear()
    from server.main import create_app

    return TestClient(create_app())


def test_export_docx(client):
    res = client.post(
        "/v1/export/docx",
        json={"session_id": "test-session", "title": "Test Export"},
    )
    assert res.status_code == 200
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    assert len(res.content) > 100


def test_export_pdf(client):
    res = client.post(
        "/v1/export/pdf",
        json={"session_id": "test-session", "title": "Test Export"},
    )
    # pandoc 可能未装 TeX，允许 200 或 503
    assert res.status_code in (200, 503)
    if res.status_code == 200:
        assert res.content[:4] == b"%PDF"


def test_docx_ingest_roundtrip():
    from docx import Document
    from server.memory.rag.doc_ingest import extract_text_from_docx

    buf = io.BytesIO()
    doc = Document()
    doc.add_heading("Test Paper", level=1)
    doc.add_paragraph("Loss landscape analysis.")
    doc.save(buf)
    text = extract_text_from_docx(buf.getvalue())
    assert "Test Paper" in text
    assert "Loss landscape" in text


def test_upload_docx_file(client, monkeypatch):
    monkeypatch.setenv("ENABLE_RAG", "true")
    from server.config import get_settings

    get_settings.cache_clear()

    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Neural network training dynamics.")
    doc.save(buf)

    res = client.post(
        "/v1/documents/upload",
        data={"session_id": "docx-test"},
        files={"file": ("paper.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 200
    data = res.json()
    assert data["document"]["title"] == "paper.docx"
